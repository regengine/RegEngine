"""POST /api/v1/ingest/csv accepts Excel .xlsx workbooks.

The workbook is converted to CSV text and then goes through the same
parse_csv_rows / ingest_events path as a CSV upload, so the same rows must
produce the same events either way.
"""

from __future__ import annotations

import csv
import io
import sys
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

_ingestion_dir = Path(__file__).resolve().parent.parent
if str(_ingestion_dir) not in sys.path:
    sys.path.insert(0, str(_ingestion_dir))
_services_dir = _ingestion_dir.parent
if str(_services_dir) not in sys.path:
    sys.path.insert(0, str(_services_dir))

import pytest
from app import csv_templates as ct
from app.authz import IngestionPrincipal, get_ingestion_principal
from app.csv_templates import MAX_CSV_ROWS, router
from app.shared.csv_safety import sanitize_cell
from app.shared.xlsx_rows import XLSX_CONTENT_TYPE, cell_to_text
from app.webhook_compat import _verify_api_key
from app.webhook_models import EventResult, IngestResponse
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

HEADER = [
    "cte_type",
    "traceability_lot_code",
    "product_description",
    "quantity",
    "unit_of_measure",
    "harvest_date",
    "location_name",
    "reference_document",
]
CSV_ROWS = [
    ["harvesting", "TLC-1", "Kale", "10", "lbs", "2026-04-17", "Farm A", "REF-1"],
    ["harvesting", "TLC-2", "Romaine", "12.5", "kg", "2026-04-18", "Farm B", "REF-2"],
    ["harvesting", "", "Kale", "10", "lbs", "2026-04-17", "Farm A", "REF-3"],  # no lot code: rejected
]
# The same rows with native Excel types: numbers and real date cells.
XLSX_ROWS = [
    ["harvesting", "TLC-1", "Kale", 10, "lbs", datetime(2026, 4, 17), "Farm A", "REF-1"],
    ["harvesting", "TLC-2", "Romaine", 12.5, "kg", date(2026, 4, 18), "Farm B", "REF-2"],
    ["harvesting", None, "Kale", 10.0, "lbs", datetime(2026, 4, 17), "Farm A", "REF-3"],
]
XLSX_NAME = "log.xlsx"


def _csv_bytes(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADER)
    writer.writerows(rows)
    return buf.getvalue().encode()


def _xlsx_bytes(sheets: dict[str, list[list[Any]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def captured(monkeypatch) -> dict:
    seen: dict = {}

    async def _fake(payload, **kwargs):
        seen["payload"] = payload
        seen["kwargs"] = kwargs
        return IngestResponse(
            accepted=len(payload.events),
            rejected=0,
            total=len(payload.events),
            events=[
                EventResult(traceability_lot_code=e.traceability_lot_code, cte_type=e.cte_type.value, status="accepted")
                for e in payload.events
            ],
            ingestion_run_id="run-1",
        )

    monkeypatch.setattr(ct, "ingest_events", _fake)
    return seen


@pytest.fixture
def client() -> TestClient:
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[_verify_api_key] = lambda: None
    app.dependency_overrides[get_ingestion_principal] = lambda: IngestionPrincipal(
        key_id="key-1", scopes=["*"], tenant_id="t1"
    )
    return TestClient(app)


def _post(client: TestClient, name: str, body: bytes, content_type: str, **data: str):
    return client.post("/api/v1/ingest/csv", files={"file": (name, body, content_type)}, data=data)


def _events(captured: dict) -> list[dict]:
    return [e.model_dump() for e in captured["payload"].events]


class TestXlsxMatchesCsv:
    def test_same_rows_same_events_and_counts(self, client, captured):
        csv_resp = _post(client, "log.csv", _csv_bytes(CSV_ROWS), "text/csv")
        assert csv_resp.status_code == 200, csv_resp.text
        csv_events = _events(captured)
        csv_receipt = captured["kwargs"]["source_file"]

        body = _xlsx_bytes({"Log": [HEADER, *XLSX_ROWS]})
        xlsx_resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE)
        assert xlsx_resp.status_code == 200, xlsx_resp.text
        xlsx_receipt = captured["kwargs"]["source_file"]

        assert _events(captured) == csv_events
        for key in ("accepted", "rejected", "total"):
            assert xlsx_resp.json()[key] == csv_resp.json()[key]
        assert xlsx_resp.json()["accepted"] == 2 and xlsx_resp.json()["rejected"] == 1
        assert xlsx_receipt.record_count == csv_receipt.record_count == 3
        assert len(xlsx_receipt.parse_errors) == 1
        # The receipt keeps the original workbook bytes, and the run records the format.
        assert xlsx_receipt.raw_bytes == body
        assert xlsx_receipt.filename == XLSX_NAME
        assert captured["payload"].source == "xlsx_upload"

    def test_detected_by_content_type_without_xlsx_name(self, client, captured):
        body = _xlsx_bytes({"Log": [HEADER, XLSX_ROWS[0]]})
        resp = _post(client, "upload", body, XLSX_CONTENT_TYPE)
        assert resp.status_code == 200, resp.text
        assert resp.json()["accepted"] == 1

    def test_explicit_source_is_kept(self, client, captured):
        body = _xlsx_bytes({"Log": [HEADER, XLSX_ROWS[0]]})
        resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE, source="erp_export")
        assert resp.status_code == 200, resp.text
        assert captured["payload"].source == "erp_export"

    def test_empty_rows_skipped_and_header_is_first_non_empty_row(self, client, captured):
        body = _xlsx_bytes({"Log": [[], [None, None], HEADER, [], XLSX_ROWS[0], [None] * 8, XLSX_ROWS[1]]})
        resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE)
        assert resp.status_code == 200, resp.text
        assert resp.json()["accepted"] == 2 and resp.json()["rejected"] == 0


class TestCellConversion:
    def test_dates_iso_and_whole_numbers_plain(self, client, captured):
        body = _xlsx_bytes({"Log": [HEADER, XLSX_ROWS[0]]})
        resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE)
        assert resp.status_code == 200, resp.text
        (event,) = captured["payload"].events
        assert event.kdes["harvest_date"] == "2026-04-17"
        assert event.quantity == 10

    def test_cell_to_text(self):
        assert cell_to_text(datetime(2026, 4, 17)) == "2026-04-17"
        assert cell_to_text(datetime(2026, 4, 17, 8, 30)) == "2026-04-17T08:30:00"
        assert cell_to_text(date(2026, 4, 17)) == "2026-04-17"
        assert cell_to_text(10) == "10"
        assert cell_to_text(10.0) == "10"
        assert cell_to_text(12.5) == "12.5"
        assert cell_to_text(None) == ""
        assert cell_to_text(True) == "TRUE"


class TestSheetSelection:
    def test_sheet_name_selects_second_sheet(self, client, captured):
        body = _xlsx_bytes({"Notes": [["just", "notes"]], "Log": [HEADER, XLSX_ROWS[0], XLSX_ROWS[1]]})
        resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE, sheet_name="Log")
        assert resp.status_code == 200, resp.text
        assert [e.traceability_lot_code for e in captured["payload"].events] == ["TLC-1", "TLC-2"]

    def test_unknown_sheet_name_returns_400(self, client, captured):
        body = _xlsx_bytes({"Log": [HEADER, XLSX_ROWS[0]]})
        resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE, sheet_name="Missing")
        assert resp.status_code == 400
        assert "Missing" in resp.json()["detail"] and "Log" in resp.json()["detail"]
        assert "payload" not in captured


class TestRejectedFiles:
    @pytest.mark.parametrize(
        ("name", "content_type"),
        [
            ("log.xls", "application/vnd.ms-excel"),
            ("log.xlsm", "application/vnd.ms-excel.sheet.macroEnabled.12"),
            ("upload", "application/vnd.ms-excel.sheet.macroEnabled.12"),
        ],
    )
    def test_legacy_and_macro_excel_return_415(self, client, captured, name, content_type):
        resp = _post(client, name, b"\xd0\xcf\x11\xe0 not really excel", content_type)
        assert resp.status_code == 415
        assert ".xlsx" in resp.json()["detail"] and "CSV" in resp.json()["detail"]
        assert "payload" not in captured

    def test_xlsx_name_but_not_zip_returns_400(self, client, captured):
        resp = _post(client, XLSX_NAME, _csv_bytes(CSV_ROWS), XLSX_CONTENT_TYPE)
        assert resp.status_code == 400
        assert "not a valid .xlsx workbook" in resp.json()["detail"]
        assert "payload" not in captured

    def test_zip_that_is_not_a_workbook_returns_400(self, client, captured):
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("hello.txt", "hi")
        resp = _post(client, XLSX_NAME, buf.getvalue(), XLSX_CONTENT_TYPE)
        assert resp.status_code == 400
        assert "not a valid .xlsx workbook" in resp.json()["detail"]

    def test_macro_workbook_renamed_to_xlsx_returns_415(self, client, captured):
        import zipfile

        src = zipfile.ZipFile(io.BytesIO(_xlsx_bytes({"Log": [HEADER, XLSX_ROWS[0]]})))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for info in src.infolist():
                zf.writestr(info, src.read(info))
            zf.writestr("xl/vbaProject.bin", b"macro")
        resp = _post(client, XLSX_NAME, buf.getvalue(), XLSX_CONTENT_TYPE)
        assert resp.status_code == 415

    def test_row_cap_applies(self, client, captured):
        rows = [[*XLSX_ROWS[0][:1], f"TLC-{i}", *XLSX_ROWS[0][2:]] for i in range(MAX_CSV_ROWS + 50)]
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [HEADER, *rows]}), XLSX_CONTENT_TYPE)
        assert resp.status_code == 200, resp.text
        assert len(captured["payload"].events) == MAX_CSV_ROWS
        errors = captured["kwargs"]["source_file"].parse_errors
        assert any(f"maximum row limit of {MAX_CSV_ROWS}" in err for err in errors)


class TestFormulaSafety:
    PAYLOAD = '=HYPERLINK("http://evil.example","click")'

    def test_formula_without_cached_value_warns_with_cell_and_header(self, client, captured):
        row = list(XLSX_ROWS[0])
        row[7] = "=1+1"
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [HEADER, row]}), XLSX_CONTENT_TYPE)

        assert resp.status_code == 200, resp.text
        assert resp.json()["accepted"] == 1
        warnings = resp.json()["kde_warnings"]
        assert any("Log!H2" in warning and "reference_document" in warning for warning in warnings)
        assert resp.json()["kde_warning_count"] == len(warnings)

    def test_plain_values_have_no_formula_warning(self, client, captured):
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [HEADER, XLSX_ROWS[0]]}), XLSX_CONTENT_TYPE)

        assert resp.status_code == 200, resp.text
        assert not any("formula" in warning.lower() for warning in resp.json()["kde_warnings"])

    def test_formula_with_cached_value_has_no_formula_warning(self, client, captured):
        row = list(XLSX_ROWS[0])
        row[7] = "=1+1"
        body = _xlsx_bytes({"Log": [HEADER, row]})
        cached = io.BytesIO()
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(io.BytesIO(body)) as source, zipfile.ZipFile(cached, "w") as output:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    cell = root.find(".//s:c[@r='H2']", namespace)
                    assert cell is not None
                    value = cell.find("s:v", namespace)
                    assert value is not None
                    value.text = "2"
                    data = ET.tostring(root, encoding="utf-8")
                output.writestr(item, data)

        resp = _post(client, XLSX_NAME, cached.getvalue(), XLSX_CONTENT_TYPE)

        assert resp.status_code == 200, resp.text
        assert not any("formula" in warning.lower() for warning in resp.json()["kde_warnings"])
        assert captured["payload"].events[0].kdes["reference_document"] == "2"

    def test_required_formula_without_cache_is_reported_on_invalid_rows(self, client, captured):
        row = list(XLSX_ROWS[0])
        row[3] = "=5*2"
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [HEADER, row]}), XLSX_CONTENT_TYPE)

        assert resp.status_code == 400
        detail = resp.json()["detail"]
        assert any("Log!D2" in warning and "quantity" in warning for warning in detail["kde_warnings"])
        assert any("quantity is required" in error for error in detail["errors"])
        assert "payload" not in captured

    def test_formula_only_sheet_reports_warning_without_data_rows(self, client, captured):
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [["=1+1"]]}), XLSX_CONTENT_TYPE)

        assert resp.status_code == 400
        detail = resp.json()["detail"]
        assert detail["message"] == "CSV contained no data rows"
        assert any("Log!A1" in warning for warning in detail["kde_warnings"])
        assert "payload" not in captured

    def test_formula_warning_lists_at_most_twenty_cells(self, client, captured):
        header = [*HEADER, *(f"extra_{index}" for index in range(22))]
        row = [*XLSX_ROWS[0], *(["=1+1"] * 22)]
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [header, row]}), XLSX_CONTENT_TYPE)

        assert resp.status_code == 200, resp.text
        formula_warnings = [warning for warning in resp.json()["kde_warnings"] if "formula" in warning.lower()]
        assert len(formula_warnings) == 1
        assert "Log!AB2" in formula_warnings[0]
        assert "Log!AC2" not in formula_warnings[0]
        assert "and 2 more" in formula_warnings[0]

    def test_formula_on_unselected_sheet_does_not_warn(self, client, captured):
        other_row = list(XLSX_ROWS[0])
        other_row[7] = "=1+1"
        body = _xlsx_bytes({"Other": [HEADER, other_row], "Log": [HEADER, XLSX_ROWS[0]]})
        resp = _post(client, XLSX_NAME, body, XLSX_CONTENT_TYPE, sheet_name="Log")

        assert resp.status_code == 200, resp.text
        assert not any("formula" in warning.lower() for warning in resp.json()["kde_warnings"])

    def test_formula_text_matches_csv_and_is_neutralized(self, client, captured):
        csv_row = ["harvesting", "TLC-1", self.PAYLOAD, "10", "lbs", "2026-04-17", "Farm A", "REF-1"]
        assert _post(client, "log.csv", _csv_bytes([csv_row]), "text/csv").status_code == 200
        csv_events = _events(captured)

        wb = Workbook()
        ws = wb.active
        ws.append(HEADER)
        ws.append(XLSX_ROWS[0])
        cell = ws.cell(row=2, column=3)
        cell.value = self.PAYLOAD
        cell.data_type = "s"  # text that looks like a formula, as Excel stores '=...
        buf = io.BytesIO()
        wb.save(buf)
        resp = _post(client, XLSX_NAME, buf.getvalue(), XLSX_CONTENT_TYPE)
        assert resp.status_code == 200, resp.text

        assert _events(captured) == csv_events
        value = captured["payload"].events[0].product_description
        assert value == self.PAYLOAD
        assert sanitize_cell(value).startswith("'=")

    def test_live_formula_is_never_carried_as_text(self, client, captured):
        row = list(XLSX_ROWS[0])
        # A real formula cell. openpyxl saves no computed value, so data_only
        # reads it as empty; Excel would supply its last computed result.
        row[7] = self.PAYLOAD
        resp = _post(client, XLSX_NAME, _xlsx_bytes({"Log": [HEADER, row]}), XLSX_CONTENT_TYPE)
        assert resp.status_code == 200, resp.text
        assert "HYPERLINK" not in resp.text
        for event in captured["payload"].events:
            assert "HYPERLINK" not in str(event.model_dump())


def test_csv_sent_as_ms_excel_content_type_still_reads_as_csv(client, captured):
    # Windows browsers label .csv uploads application/vnd.ms-excel.
    resp = _post(client, "log.csv", _csv_bytes(CSV_ROWS), "application/vnd.ms-excel")
    assert resp.status_code == 200, resp.text
    assert resp.json()["accepted"] == 2
    assert captured["payload"].source == "csv_upload"
