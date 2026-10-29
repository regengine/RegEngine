"""Read-only diagnostic contract using the existing synthetic harvest fixture."""

from __future__ import annotations

import csv
import io
import json
import logging
import re
import zipfile
from pathlib import Path

import httpx2
import pytest
from app import authz, csv_templates as ct, pilot_diagnostic as diagnostic
from app.authz import IngestionPrincipal, get_ingestion_principal
from app.shared import xlsx_rows
from app.shared.xlsx_rows import XLSX_CONTENT_TYPE
from app.webhook_compat import _verify_api_key
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from shared.jev import JevDecider, JevSettings, review_policy

FIXTURE = Path(__file__).parent / "fixtures" / "01_harvesting.csv"
PATH = "/api/v1/ingest/csv/diagnostic"


def fixture_rows():
    reader = csv.DictReader(io.StringIO(FIXTURE.read_text()))
    return list(reader.fieldnames), list(reader)


def csv_bytes(header, rows):
    out = io.StringIO()
    writer = csv.DictWriter(out, header, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode()


def xlsx_bytes(header, rows, *, padding=False):
    workbook = Workbook()
    sheet = workbook.active
    if padding:
        sheet.append([])
    sheet.append(header)
    for row in rows:
        if padding:
            sheet.append([])
        sheet.append([row.get(field) for field in header])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


@pytest.mark.parametrize("record_count,declared_last_row,status", [(3, 2, 200), (501, 501, 413), (3, 30000, 200)])
def test_xlsx_preview_uses_stored_rows_instead_of_declared_dimensions(
    client, monkeypatch, record_count, declared_last_row, status
):
    header, rows = fixture_rows()
    original = xlsx_bytes(header, [rows[0]] * record_count)
    altered = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(altered, "w") as destination:
        for entry in source.infolist():
            contents = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                contents, changed = re.subn(
                    rb'<dimension ref="[^"]+"',
                    f'<dimension ref="A1:J{declared_last_row}"'.encode(),
                    contents,
                    count=1,
                )
                assert changed == 1
            destination.writestr(entry, contents)

    if status == 413:

        def forbidden(*args, **kwargs):
            pytest.fail("Oversized workbook must be rejected before canonical parsing or Jev")

        monkeypatch.setattr(ct, "parse_csv_rows", forbidden)
        monkeypatch.setattr(review_policy, "review_pilot_diagnostic", forbidden)
    response = post(client, altered.getvalue(), xlsx=True, semantic_review="false")
    assert response.status_code == status
    if status == 200:
        assert response.json()["parsed_record_count"] == record_count
        assert response.json()["valid_event_count"] == record_count


@pytest.fixture
def app(monkeypatch):
    app = FastAPI()
    app.include_router(ct.router)
    app.dependency_overrides[_verify_api_key] = lambda: None
    app.dependency_overrides[get_ingestion_principal] = lambda: IngestionPrincipal(
        key_id="synthetic-key", scopes=["webhooks.ingest"], tenant_id="synthetic-tenant"
    )
    monkeypatch.setattr(authz, "consume_tenant_rate_limit", lambda **kwargs: (True, 100))

    def forbidden(*args, **kwargs):
        pytest.fail("A diagnostic preview must not ingest or finalize persisted events")

    monkeypatch.setattr(ct, "ingest_events", forbidden)
    monkeypatch.setattr(ct, "finalize_row_dispositions_from_ingest", forbidden)
    return app


@pytest.fixture
def client(app):
    return TestClient(app)


def post(client, body=None, *, xlsx=False, **data):
    return client.post(
        PATH,
        files={
            "file": (
                "synthetic-private-filename.xlsx" if xlsx else "synthetic-private-filename.csv",
                body if body is not None else FIXTURE.read_bytes(),
                XLSX_CONTENT_TYPE if xlsx else "text/csv",
            )
        },
        data={"cte_type": "harvesting", **data},
    )


def test_default_is_read_only_and_replay_has_no_side_effects(client, monkeypatch):
    def forbidden():
        pytest.fail("No model work without explicit opt-in")

    monkeypatch.setattr(review_policy, "get_jev_decider", forbidden)
    first = post(client)
    second = post(client)
    assert first.status_code == second.status_code == 200
    assert first.json() == second.json()
    body = first.json()
    assert body["status"] == "preview"
    assert body["parsed_record_count"] == body["valid_event_count"] == 3
    assert body["persisted_event_count"] == 0
    assert body["automatic_assignment_enabled"] is False
    assert body["regulatory_completeness"] == "not_determined"
    assert body["semantic_review"]["attempted_cluster_count"] == 0


@pytest.mark.parametrize(
    "case,expected",
    [
        ("missing_auth", 401),
        ("missing_permission", 403),
        ("form_tenant", 403),
        ("header_tenant", 403),
        ("conflicting_tenant", 400),
    ],
)
def test_authorization_and_tenant_binding_precede_read_parse_and_model(app, monkeypatch, case, expected):
    def forbidden(*args, **kwargs):
        pytest.fail("Rejected callers must not read, parse, or invoke Jev")

    monkeypatch.setattr(diagnostic, "read_upload_with_limit", forbidden)
    monkeypatch.setattr(diagnostic, "_parse_preview", forbidden)
    monkeypatch.setattr(review_policy, "get_jev_decider", forbidden)
    data = {"semantic_review": "true", "cte_type": "harvesting"}
    headers = {}
    if case == "missing_auth":
        app.dependency_overrides.pop(get_ingestion_principal)
        monkeypatch.setenv("REGENGINE_ENV", "production")
        monkeypatch.setattr(authz, "_check_auth_rate_limit", lambda *a: None)
        monkeypatch.setattr(authz, "_record_auth_failure", lambda *a: None)
    elif case == "missing_permission":
        app.dependency_overrides[get_ingestion_principal] = lambda: IngestionPrincipal(
            key_id="synthetic-key", scopes=["webhooks.read"], tenant_id="synthetic-tenant"
        )
    elif case == "form_tenant":
        data["tenant_id"] = "other-tenant"
    elif case == "header_tenant":
        headers["X-Tenant-ID"] = "other-tenant"
    else:
        data["tenant_id"] = "other-tenant"
        headers["X-Tenant-ID"] = "synthetic-tenant"
    response = TestClient(app).post(
        PATH, files={"file": ("fixture.csv", FIXTURE.read_bytes())}, data=data, headers=headers
    )
    assert response.status_code == expected


def test_mixed_clusters_retain_exact_rows_and_counts(client):
    header, rows = fixture_rows()
    rows[0]["quantity"] = ""
    rows[1]["quantity"] = ""
    rows[2]["reference_document"] = ""
    response = post(client, csv_bytes(header, rows))
    assert response.status_code == 200, response.text
    body = response.json()
    assert (
        body["parsed_record_count"],
        body["valid_event_count"],
        body["rejected_record_count"],
        body["skipped_record_count"],
        body["unparsed_record_count"],
    ) == (3, 1, 2, 0, 0)
    quantity = next(issue for issue in body["issues"] if issue["reason"] == "missing_quantity")
    reference = next(issue for issue in body["issues"] if issue["field"] == "reference_document")
    assert quantity["row_indices"] == [2, 3] and quantity["affected_row_count"] == 2
    assert reference["row_indices"] == [4] and reference["affected_row_count"] == 1


@pytest.mark.parametrize("include_input_columns", [True, False])
def test_transformation_input_warnings_are_row_associated_and_private(
    client, monkeypatch, caplog, include_input_columns
):
    header, rows = fixture_rows()
    header[header.index("harvest_date")] = "transformation_date"
    header.extend(["cte_type", "input_lot_codes"])
    if include_input_columns:
        header.extend(["input_quantities", "input_units_of_measure"])
    for row in rows:
        row["transformation_date"] = row.pop("harvest_date")
        row["cte_type"] = "transformation"
    rows[0]["input_lot_codes"] = rows[1]["traceability_lot_code"]
    if include_input_columns:
        rows[1].update(
            input_lot_codes=rows[2]["traceability_lot_code"], input_quantities="10", input_units_of_measure="lbs"
        )
    text = csv_bytes(header, rows)
    events, errors, warnings = ct.parse_csv_rows(text.decode())
    assert len(events) == 3 and errors == []
    assert sum("input_quantities missing" in warning for warning in warnings) == 1
    states = []

    async def review(state):
        states.append(state)
        return {"status": "review", "automatic_assignment_enabled": False}

    monkeypatch.setattr(review_policy, "review_pilot_diagnostic", review)
    caplog.set_level(logging.INFO)
    response = post(client, text, cte_type="transformation", semantic_review="true")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["valid_event_count"] == 3 and body["rejected_record_count"] == 0
    missing = {issue["field"]: issue for issue in body["issues"] if issue["reason"] == "missing_input_quantity_data"}
    assert set(missing) == {"input_quantities", "input_units_of_measure"}
    for issue in missing.values():
        assert issue["disposition"] == "flagged" and issue["row_indices"] == [2]
        assert issue["affected_row_count"] == 1
    assert {state["error_cluster"]["field"] for state in states} == set(missing)
    rendered = response.text + json.dumps(states) + caplog.text
    for row in rows:
        for field in ("traceability_lot_code", "product_description", "location_name", "reference_document"):
            assert row[field] not in rendered


@pytest.mark.parametrize(
    "lot_code,reason",
    [
        ("LOT1O23", "lot_code_possible_character_swap"),
        ("LOT1I23", "lot_code_possible_character_swap"),
        ("ALPHABETICAL", "lot_code_no_digits"),
    ],
)
def test_lot_code_integrity_warnings_survive_preview_without_values(client, monkeypatch, caplog, lot_code, reason):
    header, rows = fixture_rows()
    header.append("cte_type")
    for row in rows:
        row["cte_type"] = "harvesting"
    rows[1]["traceability_lot_code"] = lot_code
    text = csv_bytes(header, rows)
    events, errors, warnings = ct.parse_csv_rows(text.decode())
    assert len(events) == 3 and errors == []
    assert any(lot_code in warning for warning in warnings)
    states = []

    async def review(state):
        states.append(state)
        return {"status": "review", "automatic_assignment_enabled": False}

    monkeypatch.setattr(review_policy, "review_pilot_diagnostic", review)
    caplog.set_level(logging.INFO)
    response = post(client, text, semantic_review="true")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["valid_event_count"] == 3 and body["rejected_record_count"] == 0
    issue = next(issue for issue in body["issues"] if issue["reason"] == reason)
    assert issue["row_indices"] == [3] and issue["affected_row_count"] == 1
    assert issue["field"] == "traceability_lot_code" and issue["disposition"] == "flagged"
    assert len(states) == 1 and states[0]["error_cluster"]["issue_reason"] == reason
    assert lot_code not in response.text + json.dumps(states) + caplog.text


def test_skipped_and_multiline_records_use_record_indices(client):
    header, rows = fixture_rows()
    rows[0]["product_description"] += "\n" + rows[0]["product_description"]
    rows[1]["quantity"] = ""
    records = [rows[0], {}, {header[0]: "# synthetic comment"}, rows[1]]
    response = post(client, csv_bytes(header, records))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["parsed_record_count"] == 4
    assert body["valid_event_count"] == body["rejected_record_count"] == 1
    assert body["skipped_record_count"] == 2 and body["unparsed_record_count"] == 0
    issues = {issue["reason"]: issue for issue in body["issues"]}
    assert issues["blank_row"]["row_indices"] == [3]
    assert issues["comment_row"]["row_indices"] == [4]
    assert issues["missing_quantity"]["row_indices"] == [5]


def test_csv_and_xlsx_share_parser_and_record_index_basis(client, monkeypatch):
    header, rows = fixture_rows()
    rows[1]["quantity"] = ""
    original = ct.parse_csv_rows
    seen = []

    def parse(*args, **kwargs):
        seen.append(True)
        return original(*args, **kwargs)

    monkeypatch.setattr(ct, "parse_csv_rows", parse)
    csv_result = post(client, csv_bytes(header, rows)).json()
    xlsx_result = post(client, xlsx_bytes(header, rows, padding=True), xlsx=True).json()
    assert len(seen) == 2
    assert csv_result["issues"] == xlsx_result["issues"]
    assert csv_result["row_index_basis"] == "csv_record_index_including_header"
    assert xlsx_result["row_index_basis"] == "converted_xlsx_record_index_including_header"
    assert next(i for i in xlsx_result["issues"] if i["reason"] == "missing_quantity")["row_indices"] == [3]


@pytest.mark.parametrize("xlsx", [False, True])
@pytest.mark.parametrize("count,expected", [(500, 200), (501, 413)])
def test_500_record_boundary(client, xlsx, count, expected):
    header, rows = fixture_rows()
    records = [rows[0]] * count
    body = xlsx_bytes(header, records) if xlsx else csv_bytes(header, records)
    response = post(client, body, xlsx=xlsx)
    assert response.status_code == expected, response.text
    if expected == 200:
        assert response.json()["parsed_record_count"] == response.json()["valid_event_count"] == count


def sparse_xlsx_bytes(initial_count, final_row):
    header, rows = fixture_rows()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for _ in range(initial_count):
        sheet.append([rows[0][field] for field in header])
    for column, field in enumerate(header, start=1):
        sheet.cell(row=final_row, column=column, value=rows[0][field])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_sparse_501st_record_beyond_scan_limit_is_rejected_before_parse_or_review(client, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("A partial workbook scan must not reach parsing or Jev")

    monkeypatch.setattr(diagnostic, "_parse_preview", forbidden)
    monkeypatch.setattr(review_policy, "get_jev_decider", forbidden)
    # Exactly 500 visible data records, followed by an unseen 501st at row 20001.
    response = post(client, sparse_xlsx_bytes(500, 20001), xlsx=True, semantic_review="true")
    assert response.status_code == 413


def test_sparse_rows_within_scan_limit_are_not_counted_as_data(client, monkeypatch):
    monkeypatch.setattr(xlsx_rows, "_MAX_XLSX_SCAN_ROWS", 6)
    response = post(client, sparse_xlsx_bytes(1, 6), xlsx=True)
    assert response.status_code == 200, response.text
    assert response.json()["parsed_record_count"] == response.json()["valid_event_count"] == 2


def test_strict_scan_limit_is_opt_in_for_existing_xlsx_helper(monkeypatch):
    monkeypatch.setattr(xlsx_rows, "_MAX_XLSX_SCAN_ROWS", 3)
    content = sparse_xlsx_bytes(1, 6)
    text, _ = xlsx_rows.xlsx_to_csv_text(content, max_data_rows=500)
    assert len(list(csv.DictReader(io.StringIO(text)))) == 1
    with pytest.raises(HTTPException) as caught:
        xlsx_rows.xlsx_to_csv_text(content, max_data_rows=500, reject_on_scan_limit=True)
    assert caught.value.status_code == 413


@pytest.mark.parametrize(
    "body,xlsx", [(b"\xff\xfe", False), (b"not a workbook", True), (b"quantity\n" + b"x" * 200000, False)]
)
def test_malformed_files_do_not_reach_model(client, monkeypatch, body, xlsx):
    def forbidden():
        pytest.fail("Malformed input must not reach Jev")

    monkeypatch.setattr(review_policy, "get_jev_decider", forbidden)
    response = post(client, body, xlsx=xlsx, semantic_review="true")
    assert response.status_code == 400


@pytest.mark.parametrize("disabled", [True, False])
def test_disabled_or_malformed_provider_preserves_findings_and_stops_review(client, monkeypatch, disabled):
    calls = []

    def respond(request):
        calls.append(request)
        return httpx2.Response(200, json={"model": "jev-1.14.0", "answers": {}})

    transport = httpx2.MockTransport(respond)
    decider = JevDecider(
        JevSettings(api_key=None if disabled else "test-key", max_retries=0),
        transport=transport,
        async_transport=transport,
    )
    monkeypatch.setattr(review_policy, "get_jev_decider", lambda: decider)
    response = post(client, semantic_review="true")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["issues"]
    review = body["issues"][0]["semantic_review"]
    assert review["status"] == "unavailable" and review["suggestions"] is None
    assert review["automatic_assignment_enabled"] is False
    assert body["semantic_review"]["attempted_cluster_count"] == 1
    assert body["semantic_review"]["reviewed_cluster_count"] == 0
    assert len(calls) == (0 if disabled else 1)


def test_aggregate_review_is_bounded_and_upload_values_never_escape(client, monkeypatch, caplog):
    header, rows = fixture_rows()
    private_header = "synthetic_private_header"
    private_cell = "synthetic_private_cell"
    header.extend([private_header, "cte_type"])
    # Independent failures retain their separate rows even after the review cap.
    mutations = [
        ("reference_document", ""),
        ("quantity", ""),
        ("quantity", private_cell),
        ("unit_of_measure", ""),
        ("harvest_date", private_cell),
        ("product_description", ""),
        ("cte_type", private_cell),
    ]
    rows = [{**rows[0], field: value, private_header: private_cell} for field, value in mutations]
    captured = []

    async def review(state):
        captured.append(state)
        return {"status": "review", "automatic_assignment_enabled": False}

    monkeypatch.setattr(review_policy, "review_pilot_diagnostic", review)
    caplog.set_level(logging.INFO)
    response = post(client, csv_bytes(header, rows), semantic_review="true")
    assert response.status_code == 200, response.text
    body = response.json()
    assert len(body["issues"]) > diagnostic.MAX_REVIEW_CLUSTERS
    assert len(captured) == diagnostic.MAX_REVIEW_CLUSTERS
    assert body["semantic_review"]["omitted_cluster_count"] == len(body["issues"]) - len(captured)
    assert body["semantic_review"]["unreviewed_cluster_count"] == len(body["issues"]) - len(captured)
    rendered = response.text + json.dumps(captured) + caplog.text
    for secret in (
        private_header,
        private_cell,
        "synthetic-private-filename",
        "synthetic-tenant",
        rows[0]["traceability_lot_code"],
        rows[0]["product_description"],
    ):
        assert secret not in rendered


def test_validation_exception_does_not_log_raw_values(client, caplog):
    header, rows = fixture_rows()
    private_value = "synthetic_private_value_" * 500
    rows[0]["product_description"] = private_value
    caplog.set_level(logging.INFO)
    response = post(client, csv_bytes(header, rows))
    assert response.status_code == 200
    assert any(issue["reason"] == "parse_exception" for issue in response.json()["issues"])
    assert private_value not in response.text + caplog.text


def test_unknown_headers_are_reported_without_echoing_them(client):
    response = post(client, b"synthetic_private_column\nsynthetic_private_cell\n")
    assert response.status_code == 200
    body = response.json()
    assert body["valid_event_count"] == 0
    assert body["rejected_record_count"] == 1
    assert "synthetic_private" not in response.text
    assert any(issue["file_level"] for issue in body["issues"])
