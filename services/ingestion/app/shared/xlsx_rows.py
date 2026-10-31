"""Turn an uploaded Excel (.xlsx) workbook into CSV text.

``POST /api/v1/ingest/csv`` accepts .xlsx as well as CSV. The workbook is
converted here into the same CSV text a user would get from "Save as CSV",
and that text then goes through ``parse_csv_rows`` and the rest of the CSV
route unchanged, so validation, hashing and persistence stay on one path.

Cell values are kept verbatim, as they are for CSV: a value such as
``=HYPERLINK(...)`` stays plain text and is neutralized by ``sanitize_cell``
wherever the data is exported back to a spreadsheet. Formula cells are read
with ``data_only=True``, so only the value Excel last computed is used and
the formula itself is never carried over. Formula cells without a saved
computed value are reported as warnings.
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Optional

from fastapi import HTTPException

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MACRO_CONTENT_TYPE = "application/vnd.ms-excel.sheet.macroenabled.12"
_UNSUPPORTED_EXCEL_SUFFIXES = {".xls", ".xlsm", ".xlsb", ".xltm"}
_ZIP_MAGIC = b"PK\x03\x04"

# Zip bomb guards. Rows and columns read are capped (the caller passes the
# CSV row cap), and the declared uncompressed size of the archive is capped
# before openpyxl opens it.
MAX_XLSX_COLUMNS = 200
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
# Empty rows are skipped but still cost a loop turn; a sheet whose rows sit
# far apart must not make us walk a million blank rows.
_MAX_XLSX_SCAN_ROWS = 20_000
_MAX_UNCACHED_FORMULA_CELLS_LISTED = 20

_UNSUPPORTED_DETAIL = (
    "Legacy .xls and macro-enabled Excel files (.xlsm, .xlsb) are not supported. "
    "Save the workbook as .xlsx (Excel Workbook) or CSV UTF-8 and upload it again."
)
_INVALID_DETAIL = (
    "File is not a valid .xlsx workbook. Password-protected or renamed files cannot be read; "
    "save it as .xlsx (Excel Workbook) or CSV UTF-8 and upload it again."
)


def is_xlsx_upload(filename: Optional[str], content_type: Optional[str], content: bytes = b"") -> bool:
    """True when the upload should be read as an .xlsx workbook.

    An ``.xlsx`` file name always selects the workbook path (bad bytes are
    then rejected by :func:`xlsx_to_csv_text`). The OOXML content type alone
    only does when the bytes are a ZIP, so a CSV sent with a wrong content
    type keeps its existing behavior.
    """
    if Path(filename or "").suffix.lower() == ".xlsx":
        return True
    declared = (content_type or "").split(";")[0].strip().lower()
    return declared == XLSX_CONTENT_TYPE and content.startswith(_ZIP_MAGIC)


def reject_unsupported_excel(filename: Optional[str], content_type: Optional[str]) -> None:
    """Raise 415 for legacy or macro-enabled Excel uploads.

    Detection is by file name and the macro-enabled content type only. The
    generic ``application/vnd.ms-excel`` type is not used, because Windows
    browsers send it for ordinary .csv files.
    """
    declared = (content_type or "").split(";")[0].strip().lower()
    if Path(filename or "").suffix.lower() in _UNSUPPORTED_EXCEL_SUFFIXES or declared == _MACRO_CONTENT_TYPE:
        raise HTTPException(status_code=415, detail=_UNSUPPORTED_DETAIL)


def cell_to_text(value: object) -> str:
    """Render one cell the way the CSV path would see it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0) and value.tzinfo is None:
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def xlsx_to_csv_text(
    content: bytes,
    *,
    sheet_name: Optional[str] = None,
    max_data_rows: int,
    reject_on_scan_limit: bool = False,
) -> tuple[str, list[str]]:
    """Convert workbook bytes to CSV text and report formulas without cached values.

    Reads the first worksheet, or ``sheet_name`` when given (400 if it does
    not exist). Fully empty rows are skipped and the first non-empty row is
    the header. At most ``max_data_rows + 1`` data rows are emitted: one past
    the cap, so ``parse_csv_rows`` still reports that the file is over its
    row limit. ``reject_on_scan_limit`` rejects a partial scan of a sparse
    worksheet instead of returning an incomplete converted file. This limit
    counts worksheet rows, including blanks, separately from data records.
    """
    if not content.startswith(_ZIP_MAGIC):
        raise HTTPException(status_code=400, detail=_INVALID_DETAIL)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=400, detail=_INVALID_DETAIL) from exc
    names = {info.filename for info in infos}
    if "xl/vbaProject.bin" in names:
        raise HTTPException(status_code=415, detail=_UNSUPPORTED_DETAIL)
    if sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise HTTPException(status_code=400, detail="Excel workbook expands to too much data to read safely")

    from openpyxl import load_workbook

    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        formula_workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises many types for malformed packages
        if workbook is not None:
            workbook.close()
        raise HTTPException(status_code=400, detail=_INVALID_DETAIL) from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise HTTPException(
                    status_code=400,
                    detail=f"Sheet {sheet_name!r} not found. Sheets in this workbook: {', '.join(workbook.sheetnames)}",
                )
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.worksheets[0]
        formula_sheet = formula_workbook[sheet.title]
        if reject_on_scan_limit:
            # Declared dimensions can omit stored rows. Strict previews must
            # inspect the actual records, subject to the bounds below.
            sheet.reset_dimensions()
            formula_sheet.reset_dimensions()

        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\n")
        emitted = 0
        headers: Optional[list[str]] = None
        uncached_formulas: list[str] = []
        uncached_count = 0
        try:
            rows = sheet.iter_rows(max_col=MAX_XLSX_COLUMNS + 1)
            formula_rows = formula_sheet.iter_rows(max_col=MAX_XLSX_COLUMNS + 1)
            for scanned, (row, formula_row) in enumerate(zip(rows, formula_rows), start=1):
                if scanned > _MAX_XLSX_SCAN_ROWS:
                    if reject_on_scan_limit:
                        raise HTTPException(status_code=413, detail="Excel worksheet exceeds the safe row scan limit")
                    break
                cells = [cell_to_text(cell.value) for cell in row]
                while cells and not cells[-1].strip():
                    cells.pop()
                if cells and headers is None:
                    headers = cells
                for column, (cached_cell, formula_cell) in enumerate(zip(row, formula_row), start=1):
                    formula_value = formula_cell.value
                    if cached_cell.value is None and (
                        formula_cell.data_type == "f"
                        or isinstance(formula_value, str)
                        and formula_value.startswith("=")
                    ):
                        uncached_count += 1
                        if len(uncached_formulas) < _MAX_UNCACHED_FORMULA_CELLS_LISTED:
                            location = f"{sheet.title}!{formula_cell.coordinate}"
                            if headers and column <= len(headers) and headers[column - 1]:
                                location += f" ({headers[column - 1][:80]})"
                            uncached_formulas.append(location)
                if not cells:
                    continue
                if len(cells) > MAX_XLSX_COLUMNS:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Excel sheet has more than {MAX_XLSX_COLUMNS} columns",
                    )
                writer.writerow(cells)
                emitted += 1
                if emitted > max_data_rows + 1:
                    break
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail=_INVALID_DETAIL) from exc
        warnings = []
        if uncached_count:
            extra = uncached_count - len(uncached_formulas)
            suffix = f", and {extra} more" if extra else ""
            warnings.append(
                "Excel formula cells have no cached values: "
                + ", ".join(uncached_formulas)
                + suffix
                + ". Recalculate and save the workbook, then upload it again."
            )
        return out.getvalue(), warnings
    finally:
        workbook.close()
        formula_workbook.close()
